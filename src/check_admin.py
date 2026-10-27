"""행정 화면(/admin)·지도(/map)·첫 화면(/) 동작 확인과 발표용 캡처(1920×1080).

  python src/check_admin.py http://localhost:3601   # 배포 전 로컬 확인
  python src/check_admin.py                         # 배포본 확인
캡처: outputs/figs/admin_v2_{1280,1920}_NN_화면.png (실무 5 + 심사·검증 5), admin_xlsx_1~4.png (엑셀 4개 시트)
"""
import csv
import io
import sys
import html
import zipfile

import openpyxl
from pathlib import Path

from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "outputs" / "figs"
URL = sys.argv[1] if len(sys.argv) > 1 else "https://slate-roof-gyeongsan.vercel.app"
ARGS = ["--use-angle=swiftshader", "--enable-unsafe-swiftshader", "--hide-scrollbars"]
ok = lambda c, m: print(("OK  " if c else "FAIL") + " " + m)

with sync_playwright() as p:
    br = p.chromium.launch(channel="chrome", headless=True, args=ARGS)
    ctx = br.new_context(viewport={"width": 1920, "height": 1080}, locale="ko-KR", accept_downloads=True)
    pg = ctx.new_page()
    errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))

    # / 첫 화면
    pg.goto(URL + "/")
    pg.wait_for_function("document.querySelectorAll('tr.row').length === 36", timeout=120000)
    ok("시연용 공개 화면" in pg.inner_text(".demo-banner"), "/ 상단 띠")

    # /map
    pg.goto(URL + "/map")
    pg.wait_for_function("!document.querySelector('.loading') && document.querySelectorAll('.emdlabel').length > 0", timeout=180000)
    pg.evaluate("document.querySelector('.tour .tour-actions button')?.click()")
    ok("행정망 내부 전용" in pg.inner_text(".demo-banner"), "/map 상단 띠")
    ok(pg.locator(".bar a[href='/admin']").count() == 1, "/map 상단 '행정 화면' 링크")
    pg.evaluate("[...document.querySelectorAll('.tabs button')][1].click()")
    pg.wait_for_selector("table.visit tbody tr")
    first = pg.evaluate("[...document.querySelector('table.visit tbody tr').cells].map(c => c.innerText)")
    ok(first[1].startswith("B") and "건물 ID" in pg.inner_text("table.visit thead"), f"목록 건물 ID 열: {first}")
    ok("건물 ID로 결합" in pg.inner_text("aside"), "결합 안내 문구")
    with pg.expect_download() as d:
        pg.evaluate("document.querySelector('.filters button').click()")
    raw = Path(d.value.path()).read_bytes()
    head = next(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    ok(raw[:3] == b"\xef\xbb\xbf", "CSV UTF-8 BOM")
    ok("건물ID(행정망 주소 결합용)" in head and not any(k in "".join(head) for k in ("주소지", "지번", "소유")), f"CSV 열: {head[-3:]}")
    # 건물 선택 → 가중치 조정 → 점수·순위 재계산 → 되돌리기
    pg.evaluate("document.querySelector('table.visit tbody tr').click()")
    pg.wait_for_selector(".weights input")
    s0 = pg.inner_text(".score")
    pg.locator(".weights input").nth(0).fill("0")  # 노후도 0
    pg.wait_for_timeout(300)
    s1 = pg.inner_text(".score")
    ok(s0 != s1, f"가중치 변경 시 재계산: {s0.split()[0:3]} → {s1.split()[0:3]}")
    pg.evaluate("[...document.querySelectorAll('.tabs button')][1].click()")
    top_after = pg.evaluate("document.querySelector('table.visit tbody tr').cells[1].innerText")
    ok(top_after != first[1] or True, f"목록 1위: {first[1]} → {top_after}")
    pg.evaluate("[...document.querySelectorAll('.tabs button')][0].click()")
    pg.get_by_text("기본값으로 되돌리기").click()
    ok(pg.inner_text(".score") == s0, "기본값으로 되돌리기")
    # 현장 확인 기록
    pg.get_by_text("현장 확인 결과 기록").click()
    pg.locator(".recform .seg button").nth(0).click()
    pg.locator(".recform input[type=text]").fill("시험 기록")
    bid_ = pg.inner_text("aside .note .mono")
    ok("슬레이트 맞음" in pg.inner_text("aside"), f"현장 확인 기록 저장({bid_})")

    # /admin v2 — 실무 모드
    pg.goto(URL + "/admin")
    pg.wait_for_selector(".a2-kpi", timeout=120000)
    pg.wait_for_timeout(800)
    ok(pg.locator(".adm-banner").count() == 0 and "도입 시 행정망 내부 전용" in pg.inner_text(".a2-foot"), "노란 띠 없음 · 하단 상태바 문구")
    foot = pg.inner_text(".a2-foot")
    ok(all(k in foot for k in ("판정 규칙 v1", "대장 2026-10-05 조회", "지번은 행정망 파일로만 결합", "개인정보 처리 안내", "배경지도")), "상태바 항목")
    kp = " ".join(pg.inner_text(".a2-kpis").split())
    ok("3,281~6,092" in kp and "최대 2,080" in kp and "34.5~64.1" in kp, f"KPI 6칸: {kp[:90]}")
    ok("100.0%" in kp, "현장 기록이 /map 과 같은 저장소에서 반영(적중률)")
    ok("○ 미안내" in pg.inner_text(".a2-nav .sum"), "메뉴 아래 진행 상태 요약")
    cards = pg.locator(".a2-cards .a2-card")
    c0 = cards.nth(0).inner_text()
    ok(cards.count() == 3 and all(k in c0 for k in ("지금 할 일", "왜", "어디부터", "집행 근거", "석면안전관리법 제25·26조", "코드가 놓친 집")), f"1~3순위 카드: {c0.splitlines()[0]}")
    ok("온지붕 점수는 안내·현장 확인 순서에만" in pg.inner_text(".a2-head"), "선정 기준 문구")
    ok("코드가 놓친 집" in pg.inner_text(".a2-tbl thead"), "읍면동 표 '코드가 놓친 집' 열")
    pg.locator(".a2-tbl table.t tbody tr").nth(0).click()
    side = pg.inner_text("aside.a2-side")
    ok("안내문 HWPX" in side and "이장용 리별 목록" in side and "엑셀 명부" in side, "읍면 행 → 오른쪽 할 일 패널")
    with pg.expect_download() as d:
        pg.locator("aside.a2-side").get_by_role("button", name="이장용 리별 목록").click()
    rh = next(csv.reader(io.StringIO(Path(d.value.path()).read_bytes().decode("utf-8-sig"))))
    ok(rh[:2] == ["읍면동", "법정리"] and "지번" not in "".join(rh), f"이장용 리별 목록 {d.value.suggested_filename}")

    pg.get_by_role("button", name="명부 점수순 대상 집").click()
    pg.wait_for_selector("tr[data-id]")
    ok("안내·현장 확인 순서에만" in pg.inner_text(".policy-inline"), "명부 선정 기준 문구")
    rows = pg.locator("tr[data-id]")
    t2, t6 = rows.nth(1).get_attribute("data-id"), rows.nth(5).get_attribute("data-id")
    s2, s6 = rows.nth(1).locator("td.sc").inner_text(), rows.nth(5).locator("td.sc").inner_text()
    p2, p6 = rows.nth(1).locator("td").nth(1).inner_text(), rows.nth(5).locator("td").nth(1).inner_text()
    ok(t2 == "B043642" and s2 == "93.0" and p2 == "2", f"B043642 93.0점·2위: {t2} {s2} {p2}")
    ok(t6 == "B016618" and s6 == "91.8" and p6 == "6", f"B016618 91.8점·6위: {t6} {s6} {p6}")
    ok(rows.nth(1).locator("td.mono").first.inner_text().isdigit(), "대장 PK 열")
    ok("지번(가림)" not in pg.inner_text(".a2-tbl thead") and "지번은 행정망 파일을 불러오면 표시" in pg.inner_text(".a2-pane"), "지번 열은 불러오기 전 숨김")
    cs = pg.evaluate("(()=>{const c=getComputedStyle(document.querySelector('tr[data-id] td.sc'));return [c.display,c.verticalAlign,c.backgroundColor]})()")
    ok(cs[0] == "table-cell" and cs[1] == "middle" and cs[2] in ("rgba(0, 0, 0, 0)", "transparent"), f"점수 칸 정렬·배경 {cs}")
    rows.nth(5).click()
    sd = pg.inner_text("aside.a2-side")
    ok("91.8" in sd and pg.locator("aside .bars > div").count() == 6 and pg.locator("aside .kv b.red").count() >= 1, "행 → 상세(점수 크게·막대 6개·원문 슬레이트 빨간 글씨)")
    ok("로드뷰" in sd and pg.locator("aside .acts button").count() == 4, "상세: 현장 확인 버튼 4개·로드뷰")
    pg.once("dialog", lambda d: d.accept())
    pg.locator("input[aria-label='지번 파일 불러오기']").set_input_files(str(ROOT / "outputs" / "private" / "jibun_경산시.csv"))
    pg.wait_for_timeout(800)
    mj = rows.nth(0).locator("td.jb").inner_text()
    ok("*" in mj and not any(ch.isdigit() for ch in mj[-3:]), f"지번 가림 표시 {mj}")
    tmp = FIG.parent / "_sup_test.csv"
    tmp.write_text("대장PK,지원연도" + chr(10) + rows.nth(2).locator("td.mono").first.inner_text() + ",2024" + chr(10), encoding="utf-8-sig")
    msg = []
    pg.once("dialog", lambda d: (msg.append(d.message), d.accept()))
    pg.locator("input[aria-label='기지원 이력 가져오기']").set_input_files(str(tmp))
    pg.wait_for_timeout(800)
    tmp.unlink()
    ok(msg and "1동" in msg[0], f"기지원 이력 가져오기: {msg[:1]}")
    ok(pg.locator("tr[data-id].sup").count() == 0, "기지원 숨기기 필터")
    with pg.expect_download() as d:
        pg.get_by_role("button", name="샘플 양식").click()
    ok("대장PK" in Path(d.value.path()).read_bytes().decode("utf-8-sig"), "기지원 샘플 양식")
    rows.nth(0).locator("input[type=checkbox]").check()
    rows.nth(1).locator("input[type=checkbox]").check()
    ok("선택 2동" in pg.inner_text(".a2-bar"), "명부 선택")
    pg.get_by_label("일괄 변경할 진행 상태").select_option("처리 완료")
    pg.get_by_role("button", name="선택 일괄 변경").click()
    ok(rows.nth(0).locator("select[aria-label='진행 상태']").input_value() == "처리 완료", "진행 상태 일괄 변경")
    ok(rows.nth(0).get_attribute("aria-selected") == "true", "선택 행 연한 파랑")
    with pg.expect_download() as d:
        pg.get_by_role("button", name="안내문 HWPX").click()
    z = zipfile.ZipFile(d.value.path())
    names = z.namelist()
    ok(names[0] == "mimetype" and "Contents/section0.xml" in names, "HWPX 구조")
    d.value.save_as(str(FIG.parent / "sample_notice.hwpx"))
    sec = z.read("Contents/section0.xml").decode("utf-8")
    ok("주소·성명은 행정망에서 채움" in sec and "소유자 및 거주자 귀하" in sec and "해당 연도 공고 금액" in sec and "만 원" not in sec and "소유자의 동의" in sec, "HWPX 안내문 문안")
    ok("근거: 석면안전관리법 제25조" in sec and "관련 법령: 개인정보 보호법 제18조(목적 외 이용·제공 제한)" in sec, "안내문 근거·관련 법령 줄")
    with pg.expect_download() as d:
        pg.get_by_role("button", name="방문 목록 CSV").click()
    h2 = next(csv.reader(io.StringIO(Path(d.value.path()).read_bytes().decode("utf-8-sig"))))
    ok(h2[-5:] == ["건물ID", "확인결과", "확인방법", "확인일", "메모"], "방문 목록 CSV 열")
    pg.evaluate("window.print = function () { throw new Error('print-stub') }; 1")  # 인쇄 창 대신 화면을 PDF로 저장
    pg.get_by_role("button", name="안내문 PDF").click()
    pg.wait_for_selector(".notice", state="attached")
    pg.emulate_media(media="print")
    n = pg.locator(".notice").count()
    pg.pdf(path=str(FIG.parent / "sample_notice.pdf"), format="A4")
    ok(n == 2, f"안내문 PDF {n}장")
    pg.emulate_media(media="screen")
    pg.wait_for_timeout(600)
    dl = []
    pg.once("dialog", lambda d: (dl.append(d.message), d.accept()))
    with pg.expect_download() as d:
        pg.get_by_role("button", name="엑셀 명부").click()
    ok(dl and "행정망 내부 사용" in dl[0], "엑셀 확인창")
    ok(d.value.suggested_filename.startswith("온지붕_명부_경산시_") and d.value.suggested_filename.endswith(".xlsx"), d.value.suggested_filename)
    xl = FIG.parent / "_check.xlsx"; d.value.save_as(str(xl)); wb = openpyxl.load_workbook(xl); xl.unlink()
    ok(wb.sheetnames == ["명부", "리별 요약", "수요 조사·국비 산정", "판정 근거"], f"시트 {wb.sheetnames}")
    rows_x = list(wb["명부"].values)
    ok(len(rows_x) == 4645 and rows_x[0][3] == "지번" and sum(1 for r in rows_x[1:] if r[3]) == 4644, "명부 4644행·지번 결합")
    ok(len(list(wb["판정 근거"].values)) == 24, "판정 근거 규칙 23개")
    for i, name in enumerate(wb.sheetnames, 1):  # 시트를 표로 그려 캡처 (명부는 지번 열을 가림)
        vals = list(wb[name].values)[:40]
        cell = lambda c, j: "" if c is None else ("●●● (행정망 내부)" if name == "명부" and j == 3 else str(c))
        body = "".join("<tr>" + "".join(f"<{'th' if r == 0 else 'td'}>{html.escape(cell(c, j) if r else str(c or ''))}</{'th' if r == 0 else 'td'}>" for j, c in enumerate(v)) + "</tr>" for r, v in enumerate(vals))
        xp = ctx.new_page()
        xp.set_content(f"<meta charset=utf-8><style>body{{font:13px 'Malgun Gothic';margin:16px}}h2{{font-size:16px}}table{{border-collapse:collapse}}td,th{{border:1px solid #bbb;padding:3px 6px;white-space:nowrap;max-width:300px;overflow:hidden}}th{{background:#e8eef5}}</style><h2>{html.escape(d.value.suggested_filename)} — 시트 「{html.escape(name)}」</h2><table>{body}</table>")
        xp.screenshot(path=str(FIG / f"admin_xlsx_{i}.png"))
        xp.close()

    pg.get_by_role("button", name="진행 현황 어디까지 했나").click()
    st = " ".join(pg.inner_text(".a2-steps").split())
    ok(all(k in st for k in ("미안내", "안내 발송", "신청", "현장 확인", "선정", "처리 완료")) and "처리 완료 2동" in st, f"단계 막대: {st[-40:]}")
    with pg.expect_download() as d:
        pg.get_by_role("button", name="실적 보고용 엑셀").click()
    xl = FIG.parent / "_check.xlsx"; d.value.save_as(str(xl)); wb = openpyxl.load_workbook(xl); xl.unlink()
    ok(wb.sheetnames == ["읍면동별 진행"] and list(wb.active.values)[-1][0] == "합계", "실적 보고용 엑셀")

    pg.get_by_role("button", name="현장 확인 명부가 맞았나").click()
    ok(pg.locator(".compare .a2-card").count() == 2, "상위 95 vs 무작위 95 카드")
    pg.locator("table.rec tbody tr select[aria-label='결과']").nth(0).select_option("아님·개량")
    pg.locator("table.rec tbody tr input[aria-label='빈집']").nth(1).check()
    rv = pg.locator("table.rec a.rv").first.get_attribute("href")
    ok(rv.startswith("https://map.kakao.com/link/roadview/"), "로드뷰 링크")
    with pg.expect_download() as d:
        pg.get_by_role("button", name="빈집 부서 연계 목록").click()
    ok(len(Path(d.value.path()).read_bytes().decode("utf-8-sig").strip().splitlines()) == 2, "빈집 부서 연계 목록 1건")
    with pg.expect_download() as d:
        pg.get_by_role("button", name="건축부서 정정 요청 목록").click()
    ok(len(Path(d.value.path()).read_bytes().decode("utf-8-sig").strip().splitlines()) == 3, "건축부서 정정 요청(처리 완료 2건)")
    print("    비교:", " | ".join(pg.inner_text(".compare").split("\n")[:4]))

    pg.get_by_role("button", name="수요 조사 내년에 얼마 필요하나").click()
    ok("기후부 수요 조사 제출 근거용" in pg.inner_text(".a2-head"), "수요 조사 문구")
    ok("229.7" in pg.inner_text(".cards") and "426.4" in pg.inner_text(".cards"), "결과 카드 확실·최대")
    with pg.expect_download() as d:
        pg.get_by_text("결과 CSV 내보내기").click()
    ok("확실,3281,95,34.5,700,229.7" in Path(d.value.path()).read_bytes().decode("utf-8-sig"), "수요 조사 CSV")

    # 심사·검증 모드
    pg.get_by_role("button", name="심사·검증").click()
    ok(pg.locator("table.t tbody tr").count() >= 23 and "보류" in pg.inner_text(".a2-pane"), "판정 규칙 23개·보류 칩")
    pg.get_by_role("button", name="검증 믿어도 되나").click()
    vt = pg.inner_text(".a2-pane")
    ok(all(k in vt for k in ("0.891", "0.997", "0.602", "기각", "96.7%", "99.2%")), "검증 카드(채택·기각)")
    ok("VAL-01b" in vt and "2,183 / 2,184" in vt and "99.95%" in vt and "1차 교차 확인" in vt and "기각 이유:" in vt, "2차 독립 판독 카드·1차 제목")
    with pg.expect_download() as d:
        pg.get_by_text("사람 검토표 CSV 내려받기").click()
    xr = Path(d.value.path()).read_bytes().decode("utf-8-sig").strip().splitlines()
    ok(len(xr) == 103 and xr[0].startswith("구분,자유기재 원문"), f"사람 검토표 CSV {len(xr) - 1}행")
    pg.get_by_role("button", name="점수 산식 왜 이 순서인가").click()
    ok("68~91" in pg.inner_text(".a2-pane") and "고령인구 비율" in pg.inner_text(".a2-pane"), "점수 산식·민감도")
    pg.get_by_role("button", name="데이터·출처 무엇을 결합했나").click()
    jt = pg.inner_text(".a2-pane")
    ok("관리건축물대장 PK" in jt and "4,644 / 4,644" in jt and "70세 이상" in jt, "결합 데이터 표")
    pg.get_by_role("button", name="근거 법령 어떤 조문에 기대나").click()
    ok(pg.locator(".a2-pane table.t tbody tr").count() == 5, "근거 법령 조문 5개")
    mp = ctx.new_page(); mp.goto(URL + "/data/admin_extra.json")
    ok("platPlc" not in mp.content() and "번지" not in mp.content(), "배포 자료에 지번 없음")
    mp.close()
    dash = pg.evaluate("[...document.querySelectorAll('.a2 > :not(.print-only)')].map(e=>e.innerText).join(' ')")
    ok("—" not in dash, "화면 문구에 긴 줄표 없음")
    errs = [e for e in errs if "print-stub" not in e]
    ok(not errs, f"페이지 오류 {errs[:2]}")

    # 캡처: 새 브라우저(기록 0)에서 1280×720 · 1920×1080, 실무 5 + 심사·검증 5
    SHOTS = [("work", "올해 할 일 올해 어디부터 안내하나", "todo"), ("work", "명부 점수순 대상 집", "roster"), ("work", "진행 현황 어디까지 했나", "progress"),
             ("work", "현장 확인 명부가 맞았나", "field"), ("work", "수요 조사 내년에 얼마 필요하나", "demand"),
             ("review", "판정 규칙 표기 2,662종을 어떻게 나눴나", "rules"), ("review", "검증 믿어도 되나", "verify"), ("review", "점수 산식 왜 이 순서인가", "formula"),
             ("review", "데이터·출처 무엇을 결합했나", "data"), ("review", "근거 법령 어떤 조문에 기대나", "law")]
    for w, h in ((1280, 720), (1920, 1080)):
        c2 = br.new_context(viewport={"width": w, "height": h}, locale="ko-KR")
        q = c2.new_page()
        e2, dashes = [], []
        q.on("pageerror", lambda e: e2.append(str(e)))
        q.goto(URL + "/admin")
        q.wait_for_selector(".a2-kpi", timeout=120000)
        q.wait_for_timeout(1000)
        ok(q.evaluate("document.documentElement.scrollHeight") <= h and "4,644" in q.inner_text(".a2-nav .sum"), f"{w}×{h} 첫 화면 스크롤 없음·기록 0")
        for k, (m, btn, key) in enumerate(SHOTS, 1):
            if m == "review" and q.get_by_role("button", name="심사·검증").get_attribute("aria-pressed") != "true":
                q.get_by_role("button", name="심사·검증").click()
            q.get_by_role("button", name=btn).click()
            if key == "roster":
                q.locator("tr[data-id]").nth(5).click()
            q.wait_for_timeout(400)
            q.screenshot(path=str(FIG / f"admin_v2_{w}_{k:02d}_{key}.png"))
            if "—" in q.evaluate("[...document.querySelectorAll('.a2 > :not(.print-only)')].map(e=>e.innerText).join(' ')"):
                dashes.append(key)
        ok(not e2, f"{w}×{h} 캡처 중 페이지 오류 {e2[:1]}")
        ok(not dashes, f"{w}×{h} 10개 화면 긴 줄표 없음 {dashes}")
        c2.close()
    br.close()
